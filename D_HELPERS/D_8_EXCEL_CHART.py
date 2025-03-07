from openpyxl.chart import BarChart, Reference
#----------------------------------------------------------------
class ExcelChartGenerator:
    """ Class for adding bar charts to Excel sheets. """
    #----------------------------------------------------------------
    @staticmethod
    def add_chart(worksheet, value_col, title):
        """ Add a bar chart to the given worksheet. """
        
        max_row = worksheet.max_row
        max_entries = min(max_row - 1, 30)  

        if max_entries == 0:
            print(f" Skipping chart for '{title}' due to insufficient data.")
            return

        chart = BarChart()
        chart.title = title
        chart.x_axis.title = "Values"
        chart.y_axis.title = "Occurrences"

        data = Reference(worksheet, min_col=2, min_row=1, max_row=max_entries + 1, max_col=2)
        categories = Reference(worksheet, min_col=1, min_row=2, max_row=max_entries + 1)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        worksheet.add_chart(chart, "E5")
